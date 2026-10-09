# -*- coding: utf-8 -*-
"""md2docx：把 Markdown 导出为 Word（支持标题/表格/代码块/列表/引用/粗体/行内码）。

用途：`docs/product/` 下的设计文档需要以 .docx 交付给评审，
      而 .docx 是**生成物**——改了 Markdown 就重新导出，不要直接编辑 Word。

用法（在仓库根执行）：
    npm run docs:docx -- docs/product/青智校园_小程序详细设计文档_V2.md \
                        docs/product/export/青智校园_小程序详细设计文档_V2.docx
    # 或直接：python scripts/docs/md2docx.py <src.md> <dst.docx>

依赖：python-docx（仅在需要导出时安装：pip install python-docx）
"""
import io
import re
import sys

from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

def set_cjk(run, name="微软雅黑"):
    run.font.name = name
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = OxmlElement('w:rFonts'); rPr.append(rFonts)
    rFonts.set(qn('w:eastAsia'), name)
    rFonts.set(qn('w:ascii'), name)
    rFonts.set(qn('w:hAnsi'), name)

def add_md_runs(par, text, base_size=10.5):
    # handle **bold** and `code`
    parts = re.split(r"(\*\*[^*]+\*\*|`[^`]+`)", text)
    for pt in parts:
        if not pt: continue
        if pt.startswith("**") and pt.endswith("**") and len(pt) > 4:
            r = par.add_run(pt[2:-2]); r.bold = True; r.font.size = Pt(base_size); set_cjk(r)
        elif pt.startswith("`") and pt.endswith("`") and len(pt) > 2:
            r = par.add_run(pt[1:-1]); r.font.name = "Consolas"; r.font.size = Pt(base_size-1)
            r.font.color.rgb = RGBColor(0xC0, 0x39, 0x2B)
        else:
            r = par.add_run(pt); r.font.size = Pt(base_size); set_cjk(r)

def convert(src, dst):
    doc = Document()
    st = doc.styles['Normal']
    st.font.name = "微软雅黑"; st.font.size = Pt(10.5)
    st.element.rPr.rFonts.set(qn('w:eastAsia'), "微软雅黑")

    lines = io.open(src, encoding="utf-8").read().split("\n")
    i = 0
    in_code = False; code_buf = []
    while i < len(lines):
        l = lines[i]
        # code fence
        if l.strip().startswith("```"):
            if not in_code:
                in_code = True; code_buf = []
            else:
                in_code = False
                if code_buf:
                    p = doc.add_paragraph()
                    p.paragraph_format.space_before = Pt(4); p.paragraph_format.space_after = Pt(6)
                    p.paragraph_format.left_indent = Inches(0.15)
                    r = p.add_run("\n".join(code_buf))
                    r.font.name = "Consolas"; r.font.size = Pt(8)
            i += 1; continue
        if in_code:
            code_buf.append(l); i += 1; continue

        s = l.strip()
        if not s:
            i += 1; continue
        if s == "---":
            i += 1; continue

        # table
        if s.startswith("|") and i+1 < len(lines) and re.match(r"^\|[\s\-:|]+\|", lines[i+1].strip()):
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                rows.append([c.strip() for c in lines[i].strip().strip("|").split("|")])
                i += 1
            rows = [r for r in rows if not all(set(c) <= set("-: ") for c in r)]
            if rows:
                ncol = max(len(r) for r in rows)
                t = doc.add_table(rows=0, cols=ncol); t.style = "Table Grid"
                t.alignment = WD_TABLE_ALIGNMENT.CENTER
                for ri, row in enumerate(rows):
                    cells = t.add_row().cells
                    for ci in range(ncol):
                        txt = row[ci] if ci < len(row) else ""
                        cell = cells[ci]
                        cell.text = ""
                        par = cell.paragraphs[0]
                        add_md_runs(par, txt, base_size=9)
                        if ri == 0:
                            for r in par.runs: r.bold = True
                doc.add_paragraph()
            continue

        # headings
        m = re.match(r"^(#{1,6})\s+(.*)$", s)
        if m:
            lvl = len(m.group(1)); txt = m.group(2)
            h = doc.add_heading(level=min(lvl, 4))
            add_md_runs(h, txt, base_size={1:18,2:15,3:13,4:11.5}.get(lvl, 11))
            i += 1; continue

        # blockquote
        if s.startswith(">"):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.3)
            add_md_runs(p, s.lstrip("> ").strip(), base_size=10)
            for r in p.runs: r.italic = True
            i += 1; continue

        # list
        if re.match(r"^[-*]\s+", s) or re.match(r"^\d+\.\s+", s):
            txt = re.sub(r"^([-*]|\d+\.)\s+", "", s)
            p = doc.add_paragraph(style="List Bullet")
            add_md_runs(p, txt, base_size=10.5)
            i += 1; continue

        p = doc.add_paragraph()
        add_md_runs(p, s, base_size=10.5)
        i += 1

    doc.save(dst)
    print("saved", dst)

def usage_error() -> None:
    sys.stderr.write('用法：python scripts/docs/md2docx.py <src.md> <dst.docx>\n')
    sys.exit(2)


def main() -> None:
    if len(sys.argv) != 3:
        usage_error()
    convert(sys.argv[1], sys.argv[2])


if __name__ == '__main__':
    main()

